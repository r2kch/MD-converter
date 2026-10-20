"""Convert Office/PDF files to GitHub-Flavored Markdown.

A pure-Python, wheel-only converter.  Each supported input format is written
to a ``.md`` file placed NEXT TO the original file, and existing output files
are never overwritten (a ``.1``/``.2``/... suffix is used instead).

Chosen toolchains (see VERDICT.md for the reasoning):
  * .docx  -> mammoth (docx -> semantic HTML) + markdownify (HTML -> GFM)
  * .pptx  -> python-pptx (direct iteration over slides/placeholders)
  * .xlsx  -> openpyxl (data_only=True for cached cell values)
  * .pdf   -> pymupdf4llm (purpose-built PDF -> markdown, text layer only)

Only raise :class:`ConversionError` (a subclass of ``Exception``) so a GUI can
catch it per-file and continue with the rest.
"""
from __future__ import annotations

import io
import os
import re
import zipfile
from pathlib import Path

import mammoth
import markdownify
import pymupdf
import pymupdf4llm
from openpyxl import load_workbook
from pptx import Presentation
from pptx.enum.shapes import PP_PLACEHOLDER


class ConversionError(Exception):
    """Raised when a single file cannot be converted to markdown.

    The message is meant to be shown to a human.
    """


SUPPORTED = {".docx", ".pptx", ".xlsx", ".pdf"}


def is_supported(path) -> bool:
    """Return True if *path*'s extension is one we can convert."""
    return Path(path).suffix.lower() in SUPPORTED


# ---------------------------------------------------------------------------
# Output-path helpers
# ---------------------------------------------------------------------------

def unique_md_path(desired: Path) -> Path:
    """Return *desired*, or *desired* with a ``.N`` suffix before ``.md``.

    ``base.md`` -> ``base.1.md`` -> ``base.2.md`` ...  at the first path that
    does not already exist.  Never returns an existing path.
    """
    desired = Path(desired)
    if not desired.exists():
        return desired
    stem, ext = os.path.splitext(desired.name)  # ext == ".md"
    n = 1
    while True:
        candidate = desired.with_name(f"{stem}.{n}{ext}")
        if not candidate.exists():
            return candidate
        n += 1


def _write_unique(desired: Path, content: str) -> Path:
    """Atomically create the output at a free path (never overwrites).

    :func:`unique_md_path` picks the candidate and ``O_CREAT|O_EXCL`` claims
    it, so two threads converting the same source file cannot clobber each
    other: the loser of the claim simply retries on the next free name.
    """
    data = content.encode("utf-8")
    while True:
        candidate = unique_md_path(desired)
        try:
            fd = os.open(candidate, os.O_CREAT | os.O_EXCL | os.O_WRONLY, 0o644)
        except FileExistsError:
            continue  # taken between the check and the open: recompute
        except OSError as exc:
            raise ConversionError(
                f"Could not write '{candidate}' (target directory not writable?): {exc}"
            ) from exc
        try:
            os.write(fd, data)
        except OSError as exc:
            os.close(fd)
            candidate.unlink(missing_ok=True)
            raise ConversionError(f"Could not write '{candidate}': {exc}") from exc
        else:
            os.close(fd)
            return candidate


# ---------------------------------------------------------------------------
# Small markdown helpers
# ---------------------------------------------------------------------------

def _clean_markdown(text: str) -> str:
    """Strip trailing whitespace, collapse blank-line runs, one trailing newline."""
    lines = [line.rstrip() for line in text.split("\n")]
    out: list[str] = []
    for line in lines:
        if line == "" and out and out[-1] == "":
            continue
        out.append(line)
    while out and out[0] == "":
        out.pop(0)
    while out and out[-1] == "":
        out.pop()
    return "\n".join(out) + "\n" if out else ""


def _gfm_escape(cell: object) -> str:
    """Render one spreadsheet cell safely inside a GFM table row."""
    if cell is None:
        return ""
    return str(cell).replace("|", "\\|").replace("\n", " ").strip()


# ---------------------------------------------------------------------------
# .docx -> markdown   (mammoth -> semantic HTML -> markdownify -> GFM)
# ---------------------------------------------------------------------------

# Map the Word "Title" style to an h1 so the document title becomes a heading
# (mammoth otherwise renders it as a plain paragraph).  This *adds* to mammoth
#'s default style map, so Heading 1..6 already work.
_DOCX_STYLE_MAP = "p[style-name='Title'] => h1"


def _convert_docx(path: Path) -> str:
    try:
        with open(path, "rb") as fh:
            result = mammoth.convert_to_html(fh, style_map=_DOCX_STYLE_MAP)
    except (OSError, zipfile.BadZipFile) as exc:
        raise ConversionError(
            f"Could not read '{path.name}': not a readable DOCX "
            f"(it may be corrupt or password-protected). Details: {exc}"
        ) from exc
    except Exception as exc:  # XML errors from encrypted/unreadable parts, etc.
        raise ConversionError(
            f"Could not convert '{path.name}' (it may be corrupt or "
            f"password-protected): {exc}"
        ) from exc

    markdown = markdownify.markdownify(
        result.value,
        heading_style="ATX",
        bullets="-",
        table_infer_header=True,
    )
    return _clean_markdown(markdown)


# ---------------------------------------------------------------------------
# .pptx -> markdown   (python-pptx, per slide: title heading + body bullets)
# ---------------------------------------------------------------------------

def _convert_pptx(path: Path) -> str:
    try:
        prs = Presentation(str(path))
        slides_md: list[str] = []
        for index, slide in enumerate(prs.slides, start=1):
            title = ""
            body: list[str] = []
            for shape in slide.shapes:
                if not shape.has_text_frame:
                    continue
                text = shape.text_frame.text.strip()
                if not text:
                    continue
                is_title = (
                    shape.is_placeholder
                    and shape.placeholder_format.type == PP_PLACEHOLDER.TITLE
                )
                if is_title:
                    if not title:
                        title = text
                elif shape.is_placeholder:
                    # Bulleted content placeholder -> markdown bullets.
                    for line in text.splitlines():
                        if line.strip():
                            body.append(f"- {line.strip()}")
                else:
                    # Any other text box -> plain paragraphs.
                    for line in text.splitlines():
                        if line.strip():
                            body.append(line.strip())

            block = [f"# {title or f'Slide {index}'}"]
            if body:
                block.append("")
                block.append("\n".join(body))
            slides_md.append("\n".join(block))
    except (OSError, zipfile.BadZipFile) as exc:
        raise ConversionError(
            f"Could not read '{path.name}': not a readable PPTX "
            f"(it may be corrupt or password-protected). Details: {exc}"
        ) from exc
    except Exception as exc:
        raise ConversionError(f"Could not convert '{path.name}': {exc}") from exc

    return _clean_markdown("\n\n".join(slides_md))


# ---------------------------------------------------------------------------
# .xlsx -> markdown   (openpyxl, one .md per sheet, GFM tables)
# ---------------------------------------------------------------------------

def _sanitize_sheet(name: str) -> str:
    """Sanitize a sheet name for use in a file name (matches NAMING.md exactly).

    Replace every character NOT in ``[A-Za-z0-9-]`` with a single ``-`` (spaces
    included), collapse any run of ``-`` into one, trim leading/trailing ``-``,
    and fall back to ``Sheet`` if the result is empty.
    """
    s = re.sub(r"[^A-Za-z0-9-]", "-", name)  # each disallowed char -> one '-'
    s = re.sub(r"-{2,}", "-", s)  # collapse consecutive runs
    s = s.strip("-")  # trim leading/trailing
    return s or "Sheet"


def _sheet_to_markdown(title: str, rows: list[list]) -> str:
    """Render one worksheet as a GFM table under an h1 with the sheet name."""
    # Drop rows that are entirely empty.
    rows = [
        r
        for r in rows
        if any(c is not None and str(c).strip() != "" for c in r)
    ]
    if not rows:
        return f"# {title}\n"

    width = max(len(r) for r in rows)
    rows = [list(r) + [""] * (width - len(r)) for r in rows]

    lines = [f"# {title}", ""]
    lines.append("| " + " | ".join(_gfm_escape(c) for c in rows[0]) + " |")
    lines.append("| " + " | ".join(["---"] * width) + " |")
    for row in rows[1:]:
        lines.append("| " + " | ".join(_gfm_escape(c) for c in row) + " |")
    return "\n".join(lines) + "\n"


def _convert_xlsx(path: Path) -> list[tuple[str, str]]:
    """Return a ``(sheet_name, markdown)`` pair for every worksheet."""
    try:
        wb = load_workbook(str(path), data_only=True)
        try:
            return [
                (ws.title, _sheet_to_markdown(ws.title, [list(r) for r in ws.iter_rows(values_only=True)]))
                for ws in wb.worksheets
            ]
        finally:
            wb.close()
    except (OSError, zipfile.BadZipFile) as exc:
        raise ConversionError(
            f"Could not read '{path.name}': not a readable XLSX "
            f"(it may be corrupt or password-protected). Details: {exc}"
        ) from exc
    except Exception as exc:
        raise ConversionError(f"Could not convert '{path.name}': {exc}") from exc


# ---------------------------------------------------------------------------
# .pdf -> markdown   (pymupdf4llm; text layer only, no OCR)
# ---------------------------------------------------------------------------

def _convert_pdf(path: Path) -> str:
    try:
        # pymupdf4llm prints a "Document parser messages" block (which on this
        # machine includes a "RapidOCR is missing" note) to pymupdf's message
        # stream.  Redirect it to a throwaway buffer so the console stays clean;
        # the returned markdown is unaffected either way.
        pymupdf.set_messages(stream=io.StringIO())
        markdown = pymupdf4llm.to_markdown(str(path))
    except (OSError, zipfile.BadZipFile) as exc:
        raise ConversionError(
            f"Could not read '{path.name}': not a readable PDF "
            f"(it may be corrupt or password-protected). Details: {exc}"
        ) from exc
    except Exception as exc:
        raise ConversionError(
            f"Could not convert '{path.name}' (it may be corrupt or "
            f"password-protected): {exc}"
        ) from exc

    if not markdown or not markdown.strip():
        raise ConversionError(
            f"PDF '{path.name}' has no extractable text layer "
            f"(scanned/image-only). OCR is not supported by this tool."
        )
    return _clean_markdown(markdown)


# ---------------------------------------------------------------------------
# Public entry point
# ---------------------------------------------------------------------------

def convert_file(path) -> list[Path]:
    """Convert *path* to markdown written next to the original file.

    Returns the list of created ``.md`` paths (one for docx/pptx/pdf; one per
    sheet for xlsx).  Raises :class:`ConversionError` on any failure.  Safe to
    call from a background thread.
    """
    src = Path(path)
    if not src.is_file():
        raise ConversionError(f"File not found or not a regular file: {src}")

    ext = src.suffix.lower()
    if ext not in SUPPORTED:
        supported = ", ".join(sorted(SUPPORTED))
        raise ConversionError(f"Unsupported file type '{ext or src.name}': supported: {supported}")

    stem = src.stem
    parent = src.parent

    if ext == ".docx":
        outputs = [(f"{stem}.md", _convert_docx(src))]
    elif ext == ".pptx":
        outputs = [(f"{stem}.md", _convert_pptx(src))]
    elif ext == ".xlsx":
        outputs = [
            (f"{stem}-{_sanitize_sheet(sheet)}.md", markdown)
            for sheet, markdown in _convert_xlsx(src)
        ]
    else:  # .pdf
        outputs = [(f"{stem}.md", _convert_pdf(src))]

    return [_write_unique(parent / name, markdown) for name, markdown in outputs]


if __name__ == "__main__":
    import sys

    for arg in sys.argv[1:]:
        try:
            created = convert_file(Path(arg))
        except ConversionError as err:
            print(f"ERROR: {err}", file=sys.stderr)
            sys.exit(1)
        for p in created:
            print(p)
